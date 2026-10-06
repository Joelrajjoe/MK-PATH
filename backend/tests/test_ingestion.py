"""Phase 4 ingestion tests - all 13 required scenarios + security units.

Fixtures are tiny (a few rows); no large test datasets are generated.
"""
import io
import zipfile
from pathlib import Path

import pytest

from app.config import settings
from app.ingestion.security import (
    SecurityError,
    check_zip_limits,
    sanitize_filename,
    validate_zip_entry,
)

CSV_OK = b"id,name,score\n1,alpha,10.5\n2,beta,20.0\n3,gamma,30.5\n"


# ---------------------------------------------------------------------------
# 1. CSV (happy path + full normalized result contract)
# ---------------------------------------------------------------------------

def test_csv_upload_and_contract(upload_fn, client):
    resp = upload_fn("people.csv", CSV_OK, mime="text/csv")
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["source_format"] == "csv"
    assert len(body["datasets"]) == 1
    ds = body["datasets"][0]
    assert ds["row_count"] == 3
    assert ds["column_count"] == 3
    assert ds["ingestion_status"] == "completed"
    assert ds["warnings"] == []
    assert ds["errors"] == []

    # Full normalized result contract via GET /api/datasets/{id}
    gid = client.get(f"/api/datasets/{ds['dataset_id']}")
    assert gid.status_code == 200
    doc = gid.json()
    for key in (
        "dataset_id", "source_file_id", "original_filename", "source_format",
        "storage_path", "table_name", "schema", "row_count", "column_count",
        "ingestion_status", "warnings", "errors", "created_at",
    ):
        assert key in doc, f"missing contract field: {key}"
    assert doc["original_filename"] == "people.csv"
    assert Path(doc["storage_path"]).exists()  # real file on E:
    assert Path(doc["normalized_path"]).exists()
    assert [c["name"] for c in doc["schema"]] == ["id", "name", "score"]

    # list endpoint
    lst = client.get("/api/datasets")
    assert lst.status_code == 200
    assert lst.json()["total"] >= 1

    # schema endpoint
    sch = client.get(f"/api/datasets/{ds['dataset_id']}/schema")
    assert sch.status_code == 200
    assert [c["name"] for c in sch.json()["columns"]] == ["id", "name", "score"]

    # preview endpoint
    prev = client.get(f"/api/datasets/{ds['dataset_id']}/preview", params={"limit": 2})
    assert prev.status_code == 200
    pdata = prev.json()
    assert pdata["row_count_returned"] == 2
    assert pdata["columns"] == ["id", "name", "score"]
    assert pdata["rows"][0]["name"] == "alpha"

    # quality before profiling -> controlled 404, never fake data
    q = client.get(f"/api/datasets/{ds['dataset_id']}/quality")
    assert q.status_code == 404
    assert q.json()["detail"]["code"] == "NOT_PROFILED"


def test_csv_bom_and_delimiter(upload_fn):
    # UTF-8 BOM + semicolon delimiter
    raw = "﻿a;b;c\n1;2;3\n".encode("utf-8")
    resp = upload_fn("eu.csv", raw, mime="text/csv")
    assert resp.status_code == 201, resp.text
    ds = resp.json()["datasets"][0]
    assert ds["column_count"] == 3
    assert ds["row_count"] == 1


def test_csv_headerless_numeric(upload_fn):
    resp = upload_fn("nums.csv", b"1,2,3\n4,5,6\n", mime="text/csv")
    assert resp.status_code == 201, resp.text
    ds = resp.json()["datasets"][0]
    assert ds["row_count"] == 2
    assert any("no header row" in w for w in ds["warnings"])


# ---------------------------------------------------------------------------
# 2. XLSX (+ workbook detection + sheet selection)
# ---------------------------------------------------------------------------

def _xlsx_bytes(sheets: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for name, (headers, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name
        ws.append(headers)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_upload_and_sheet_selection(upload_fn):
    raw = _xlsx_bytes(
        {
            "Alpha": (["name", "score"], [["a", 1], ["b", 2]]),
            "Beta": (["beta_col"], [["x"], ["y"], ["z"]]),
        }
    )
    # No sheet selected -> first sheet + warning listing the workbook sheets
    resp = upload_fn("book.xlsx", raw)
    assert resp.status_code == 201, resp.text
    body = resp.json()
    ds = body["datasets"][0]
    assert ds["row_count"] == 2
    assert any("Alpha" in w and "Beta" in w for w in ds["warnings"]), ds["warnings"]

    # Explicit sheet selection -> second sheet data
    resp2 = upload_fn("book.xlsx", raw, sheet_name="Beta")
    assert resp2.status_code == 201, resp2.text
    ds2 = resp2.json()["datasets"][0]
    assert ds2["row_count"] == 3
    # missing sheet -> clear error listing available sheets
    resp3 = upload_fn("book.xlsx", raw, sheet_name="Nope")
    assert resp3.status_code == 422
    assert "Nope" in resp3.json()["detail"]["message"]
    assert "Alpha" in resp3.json()["detail"]["message"]


def test_xls_upload(upload_fn):
    xlwt = pytest.importorskip("xlwt", reason="xlwt not installed for .xls fixtures")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("Legacy")
    ws.write(0, 0, "name")
    ws.write(0, 1, "value")
    ws.write(1, 0, "row1")
    ws.write(1, 1, 42)
    buf = io.BytesIO()
    wb.save(buf)
    resp = upload_fn("legacy.xls", buf.getvalue())
    assert resp.status_code == 201, resp.text
    ds = resp.json()["datasets"][0]
    assert ds["row_count"] == 1
    assert ds["column_count"] == 2


# ---------------------------------------------------------------------------
# 3. JSON (array + NDJSON + malformed + unsupported structure)
# ---------------------------------------------------------------------------

def test_json_array(upload_fn):
    raw = b'[{"city":"Paris","pop":2},{"city":"Rome","pop":3}]'
    resp = upload_fn("cities.json", raw, mime="application/json")
    assert resp.status_code == 201, resp.text
    ds = resp.json()["datasets"][0]
    assert ds["row_count"] == 2
    assert ds["column_count"] == 2


def test_json_ndjson(upload_fn):
    raw = b'{"a":1}\n{"a":2}\n{"a":3}\n'
    resp = upload_fn("lines.json", raw, mime="application/json")
    assert resp.status_code == 201, resp.text
    ds = resp.json()["datasets"][0]
    assert ds["row_count"] == 3
    assert any("NDJSON" in w for w in ds["warnings"])


def test_json_nested_flattened(upload_fn):
    raw = b'[{"id":1,"owner":{"name":"joel","age":30}}]'
    resp = upload_fn("nested.json", raw, mime="application/json")
    assert resp.status_code == 201, resp.text
    assert resp.json()["datasets"]  # ingested
    assert any("flattened" in w for d in resp.json()["datasets"] for w in d["warnings"])


def test_json_malformed(upload_fn):
    resp = upload_fn("bad.json", b'{"a": [1, 2', mime="application/json")
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail["code"] in {"MALFORMED_JSON", "UNSUPPORTED_JSON_STRUCTURE"}
    assert "line" in detail["message"].lower()


def test_json_unsupported_structure(upload_fn):
    resp = upload_fn("scalar.json", b"42", mime="application/json")
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "UNSUPPORTED_JSON_STRUCTURE"


# ---------------------------------------------------------------------------
# 4. Parquet (happy + malformed)
# ---------------------------------------------------------------------------

def _parquet_bytes() -> bytes:
    import pandas as pd

    df = pd.DataFrame({"k": [1, 2, 3], "v": ["x", "y", "z"]})
    buf = io.BytesIO()
    df.to_parquet(buf, index=False, engine="pyarrow")
    return buf.getvalue()


def test_parquet_upload(upload_fn):
    resp = upload_fn("data.parquet", _parquet_bytes())
    assert resp.status_code == 201, resp.text
    ds = resp.json()["datasets"][0]
    assert ds["row_count"] == 3
    assert ds["column_count"] == 2
    # normalized storage_path points at the original parquet (already columnar)


def test_parquet_malformed(upload_fn):
    resp = upload_fn("bad.parquet", b"PAR1this_is_not_a_real_parquet_file")
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "MALFORMED_PARQUET"


# ---------------------------------------------------------------------------
# 5. SQL (safe dump + unsafe rejections)
# ---------------------------------------------------------------------------

SQL_SAFE = (
    b"CREATE TABLE users (id INTEGER, name VARCHAR, score DOUBLE);\n"
    b"INSERT INTO users VALUES (1, 'alice', 9.5);\n"
    b"INSERT INTO users VALUES (2, 'bob', 8.0);\n"
    b"INSERT INTO users VALUES (3, 'carol', 7.25);\n"
)


def test_sql_safe_dump(upload_fn):
    resp = upload_fn("dump.sql", SQL_SAFE, mime="application/sql")
    assert resp.status_code == 201, resp.text
    ds = resp.json()["datasets"][0]
    assert ds["row_count"] == 3
    assert ds["column_count"] == 3


def test_sql_table_name_recorded(upload_fn, client):
    resp = upload_fn("dump.sql", SQL_SAFE, mime="application/sql")
    assert resp.status_code == 201, resp.text
    ds_id = resp.json()["datasets"][0]["dataset_id"]
    doc = client.get(f"/api/datasets/{ds_id}").json()
    assert doc["table_name"] == "users"
    prev = client.get(f"/api/datasets/{ds_id}/preview")
    assert prev.status_code == 200
    assert prev.json()["row_count_returned"] == 3
    assert prev.json()["rows"][0]["name"] == "alice"


def test_sql_unsafe_only_statements_rejected(upload_fn):
    raw = b"DROP TABLE users; DELETE FROM users; COPY users TO '/tmp/x.csv';"
    resp = upload_fn("evil.sql", raw, mime="application/sql")
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail["code"] == "SQL_NO_SUPPORTED_STATEMENTS"
    assert "DROP" in detail["message"] or "Forbidden" in detail["message"]


def test_sql_destructive_statement_not_executed(upload_fn, client):
    # Mixed: valid CREATE + destructive DROP -> ingest proceeds, DROP rejected.
    raw = (
        b"CREATE TABLE t (id INTEGER);\n"
        b"INSERT INTO t VALUES (1);\n"
        b"DROP TABLE t;\n"
    )
    resp = upload_fn("mixed.sql", raw, mime="application/sql")
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert len(body["datasets"]) == 1
    warnings = body["datasets"][0]["warnings"]
    assert any("rejected" in w for w in warnings), warnings
    # dataset still has its row -> DROP never executed
    ds_id = body["datasets"][0]["dataset_id"]
    assert client.get(f"/api/datasets/{ds_id}").json()["row_count"] == 1


def test_sql_filesystem_access_blocked(upload_fn):
    raw = (
        b"CREATE TABLE s (v VARCHAR);\n"
        b"INSERT INTO s SELECT * FROM read_csv('/etc/passwd');\n"
    )
    resp = upload_fn("sneaky.sql", raw, mime="application/sql")
    # INSERT..SELECT is rejected; CREATE alone ingests with a warning,
    # but read_csv() must never execute.
    assert resp.status_code == 201, resp.text
    warnings = resp.json()["datasets"][0]["warnings"]
    assert any("rejected" in w for w in warnings), warnings


# ---------------------------------------------------------------------------
# 6. ZIP with mixed formats
# ---------------------------------------------------------------------------

def _zip_bytes(entries: dict) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in entries.items():
            zf.writestr(name, content)
    return buf.getvalue()


def test_zip_mixed_formats(upload_fn):
    raw = _zip_bytes(
        {
            "tables/data.csv": CSV_OK,
            "tables/items.json": b'[{"sku":1},{"sku":2}]',
            "tables/report.xlsx": _xlsx_bytes(
                {"Sheet1": (["c1", "c2"], [[1, 2], [3, 4]])}
            ),
            "tables/blob.parquet": _parquet_bytes(),
            "tables/notes.png": b"\x89PNG_not_really",
            "tables/script.exe": b"MZ fake binary",
        }
    )
    resp = upload_fn("bundle.zip", raw, mime="application/zip")
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["source_format"] == "zip"

    # every supported file became a dataset (processed independently)
    formats = sorted(d["source_format"] for d in body["datasets"])
    assert formats == ["csv", "json", "parquet", "xlsx"], formats

    # unsupported files are reported separately - NEVER silently ignored
    unsupported_paths = [u["path"] for u in body["unsupported_files"]]
    assert any(p.endswith("notes.png") for p in unsupported_paths)
    assert any(p.endswith("script.exe") for p in unsupported_paths)

    # manifest exists and covers every member
    manifest_paths = [m["path"] for m in body["manifest"]]
    assert any(p.endswith("data.csv") for p in manifest_paths)
    assert any(p.endswith("notes.png") for p in manifest_paths)

    # normalized results persisted and readable
    for d in body["datasets"]:
        assert d["ingestion_status"] == "completed"
    assert len(body["datasets"]) == 4


# ---------------------------------------------------------------------------
# 7. Invalid ZIP
# ---------------------------------------------------------------------------

def test_invalid_zip_rejected(upload_fn):
    resp = upload_fn("fake.zip", b"this is definitely not a zip archive", mime="application/zip")
    assert resp.status_code == 400
    assert resp.json()["detail"]["code"] == "INVALID_ZIP"


def test_empty_zip_rejected(upload_fn):
    resp = upload_fn("empty.zip", _zip_bytes({}), mime="application/zip")
    assert resp.status_code == 422  # PK\x05\x06 signature ok, but no files
    assert resp.json()["detail"]["code"] == "EMPTY_ZIP"


# ---------------------------------------------------------------------------
# 8. ZIP Slip attempt
# ---------------------------------------------------------------------------

def test_zip_slip_blocked(upload_fn):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../evil.txt", b"pwned")
        zf.writestr("..\\..\\evil2.txt", b"pwned")
        zf.writestr("/abs/evil3.txt", b"pwned")
        zf.writestr("ok/good.csv", CSV_OK)
    resp = upload_fn("slip.zip", buf.getvalue(), mime="application/zip")

    assert resp.status_code == 201, resp.text
    body = resp.json()
    rejected = {r["entry"] for r in body["rejected_entries"]}
    assert "ok/good.csv" not in rejected
    assert len(body["datasets"]) == 1  # the safe member still ingested
    assert body["datasets"][0]["original_filename"] == "good.csv"
    codes = {r["code"] for r in body["rejected_entries"]}
    assert "PATH_TRAVERSAL" in codes
    assert "ABSOLUTE_PATH_ENTRY" in codes

    # prove nothing escaped: no evil files anywhere under temp/ or project root
    root = settings.ROOT_DIR
    assert not (root / "evil.txt").exists()
    assert not (root / "evil2.txt").exists()
    assert not (settings.TEMP_DIR / "evil.txt").exists()
    for p in settings.TEMP_DIR.rglob("evil*"):
        assert False, f"ZIP slip wrote outside extraction root: {p}"


def test_zip_slip_only_malicious(upload_fn):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("../evil.txt", b"pwned")
    resp = upload_fn("allbad.zip", buf.getvalue(), mime="application/zip")
    assert resp.status_code == 422  # nothing ingestible -> structured failure
    assert not (settings.ROOT_DIR / "evil.txt").exists()


# ---------------------------------------------------------------------------
# 9. Unsupported extension
# ---------------------------------------------------------------------------

def test_unsupported_extension(upload_fn):
    resp = upload_fn("program.exe", b"MZ\x90\x00fake")
    assert resp.status_code == 415
    assert resp.json()["detail"]["code"] == "UNSUPPORTED_EXTENSION"


def test_empty_upload_rejected(upload_fn):
    resp = upload_fn("empty.csv", b"")
    assert resp.status_code == 400
    assert resp.json()["detail"]["code"] == "EMPTY_FILE"


def test_upload_size_limit(upload_fn, monkeypatch):
    monkeypatch.setattr(settings, "MAX_UPLOAD_MB", 0)
    resp = upload_fn("big.csv", CSV_OK, mime="text/csv")
    assert resp.status_code == 413
    assert resp.json()["detail"]["code"] == "UPLOAD_TOO_LARGE"


def test_binary_content_with_text_extension(upload_fn):
    resp = upload_fn("weird.csv", b"PK\x03\x04\x00\x00binary\x00junk", mime="text/csv")
    assert resp.status_code == 415
    assert resp.json()["detail"]["code"] == "BINARY_NOT_TEXT"


# ---------------------------------------------------------------------------
# 10. Malformed CSV
# ---------------------------------------------------------------------------

def test_malformed_csv_ragged_rows(upload_fn):
    resp = upload_fn("ragged.csv", b"a,b,c\n1,2\n3,4,5\n6,7,8,9\n", mime="text/csv")
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail["code"] == "MALFORMED_CSV"
    assert "line" in detail["message"].lower()


def test_malformed_csv_unbalanced_quotes(upload_fn):
    resp = upload_fn("quotes.csv", b'a,b\n1,"unclosed\n2,3\n', mime="text/csv")
    assert resp.status_code in (201, 422)  # engine-dependent; never 500
    if resp.status_code == 422:
        assert resp.json()["detail"]["code"] == "MALFORMED_CSV"


# ---------------------------------------------------------------------------
# 11-13 covered above: malformed JSON / malformed Parquet / unsafe SQL
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# Nested ZIP depth + duplicates
# ---------------------------------------------------------------------------

def test_nested_zip_within_depth(upload_fn):
    inner = _zip_bytes({"inner/data.csv": CSV_OK})
    outer = _zip_bytes({"top.csv": CSV_OK, "nested.zip": inner})
    resp = upload_fn("outer.zip", outer, mime="application/zip")
    assert resp.status_code == 201, resp.text
    names = [d["original_filename"] for d in resp.json()["datasets"]]
    assert "top.csv" in names
    assert "data.csv" in names  # one nesting level allowed (MAX_ZIP_DEPTH=1)


def test_nested_zip_beyond_depth_reported(upload_fn):
    level2 = _zip_bytes({"deep.csv": CSV_OK})
    level1 = _zip_bytes({"l2.zip": level2})
    outer = _zip_bytes({"l1.zip": level1})
    resp = upload_fn("deep.zip", outer, mime="application/zip")
    # No ingestible files; the over-depth nested archive is REPORTED in a
    # structured error - never silently ignored.
    assert resp.status_code == 422
    detail = resp.json()["detail"]
    assert detail["code"] in {"NO_SUPPORTED_FILES", "EMPTY_ZIP"}
    assert "MAX_ZIP_DEPTH" in detail["message"]


def test_duplicate_filenames_disambiguated(upload_fn):
    raw = _zip_bytes({"a.csv": CSV_OK, "./a.csv": CSV_OK})
    resp = upload_fn("dupes.zip", raw, mime="application/zip")
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert len(body["datasets"]) == 2
    dup_entries = [r for r in body["rejected_entries"] if r["code"] == "DUPLICATE_NAME"]
    assert len(dup_entries) == 1


# ---------------------------------------------------------------------------
# Pure security unit tests (no server roundtrip)
# ---------------------------------------------------------------------------

def test_sanitize_filename_never_yields_paths():
    assert "/" not in sanitize_filename("../../etc/passwd")
    assert "\\" not in sanitize_filename("C:\\Windows\\evil.exe")
    assert sanitize_filename("...") == "unnamed"
    assert sanitize_filename("normal file (1).csv") == "normal file (1).csv"


def test_zip_limits_estimated_size():
    infos = []
    info = zipfile.ZipInfo("big.bin")
    info.file_size = 300 * 1024 * 1024  # 300MB declared
    info.compress_size = 1024
    infos.append(info)
    with pytest.raises(SecurityError) as exc:
        check_zip_limits(infos)
    assert exc.value.code == "ESTIMATED_SIZE_EXCEEDED"


def test_zip_limits_ratio_bomb():
    info = zipfile.ZipInfo("bomb.bin")
    info.file_size = 50 * 1024 * 1024
    info.compress_size = 100 * 1024  # 500:1 ratio
    with pytest.raises(SecurityError) as exc:
        check_zip_limits([info])
    assert exc.value.code == "COMPRESSION_RATIO_EXCEEDED"


def test_zip_entry_traversal_rejected():
    info = zipfile.ZipInfo("../../outside.txt")
    with pytest.raises(SecurityError) as exc:
        validate_zip_entry(info)
    assert exc.value.code == "PATH_TRAVERSAL"

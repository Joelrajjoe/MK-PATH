"""Format parsers for the Universal Data Ingestion Engine (Phase 4).

All parsing is DETERMINISTIC (spec section 7): no LLM involved anywhere in
byte parsing, schema inference, or validation.

Every parser returns normalized tables backed by Parquet on E:, plus a
schema derived from Parquet metadata (never guessed by a model).
"""
import csv
import io
import json
import re
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq

from ..config import settings


class IngestionError(Exception):
    """Structured, user-facing ingestion failure."""

    def __init__(self, code: str, message: str, http_status: int = 422):
        super().__init__(message)
        self.code = code
        self.message = message
        self.http_status = http_status


@dataclass
class TableInfo:
    """One normalized table produced from one source file (or SQL table)."""

    table_name: str
    normalized_path: str
    row_count: int
    column_count: int
    schema: List[Dict[str, Any]]
    warnings: List[str] = field(default_factory=list)


def _schema_from_parquet(path: Path) -> tuple:
    """Derive schema + row count from Parquet metadata (lazy, no full load)."""
    pf = pq.ParquetFile(path)
    schema = [
        {"name": f.name, "type": str(f.type), "nullable": bool(f.nullable)}
        for f in pf.schema_arrow
    ]
    return schema, pf.metadata.num_rows


def _df_to_table(df: pd.DataFrame, dest: Path, table_name: str) -> TableInfo:
    dest.parent.mkdir(parents=True, exist_ok=True)
    df.to_parquet(dest, index=False, engine="pyarrow")
    schema, rows = _schema_from_parquet(dest)
    return TableInfo(
        table_name=table_name,
        normalized_path=str(dest),
        row_count=rows,
        column_count=len(schema),
        schema=schema,
    )


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------

def _detect_encoding(raw: bytes, warnings: List[str]) -> str:
    sample = raw[: 256 * 1024]
    try:
        sample.decode("utf-8-sig")
        return "utf-8-sig"  # also handles a leading BOM
    except UnicodeDecodeError:
        pass
    try:
        sample.decode("utf-8")
        return "utf-8"
    except UnicodeDecodeError:
        warnings.append("File is not valid UTF-8; decoded as cp1252 (best effort).")
        return "cp1252"


def _detect_delimiter(text: str, warnings: List[str]) -> str:
    try:
        dialect = csv.Sniffer().sniff(text[:64 * 1024], delimiters=",;\t|")
        return dialect.delimiter
    except csv.Error:
        pass
    first_line = text.splitlines()[0] if text.splitlines() else ""
    counts = {d: first_line.count(d) for d in [",", ";", "\t", "|"]}
    best = max(counts, key=counts.get)
    if counts[best] == 0:
        warnings.append("Could not detect a delimiter; defaulted to comma.")
        return ","
    warnings.append(f"Delimiter sniffing fell back to first-line counting: {best!r}.")
    return best


def _row_looks_numeric(values: List[str]) -> bool:
    for v in values:
        try:
            float(v)
        except ValueError:
            return False
    return True


def _detect_header(rows: List[List[str]], warnings: List[str]) -> bool:
    """Header present unless every first-row cell parses as a number."""
    if not rows:
        return True
    first = [c.strip() for c in rows[0]]
    if all(c == "" for c in first):
        return True
    if _row_looks_numeric([c for c in first if c != ""]):
        warnings.append(
            "First row looks numeric; treated as data (no header row detected)."
        )
        return False
    return True


def parse_csv(path: Path, logical_name: str) -> TableInfo:
    warnings: List[str] = []
    raw = path.read_bytes()
    if not raw.strip():
        raise IngestionError("EMPTY_FILE", "CSV file is empty (0 usable bytes).")

    encoding = _detect_encoding(raw, warnings)
    text = raw.decode(encoding, errors="replace")

    delimiter = _detect_delimiter(text, warnings)

    sample_rows = list(csv.reader(io.StringIO(text[:64 * 1024]), delimiter=delimiter))
    has_header = _detect_header(sample_rows, warnings)

    read_kwargs: Dict[str, Any] = dict(
        sep=delimiter,
        encoding=encoding,
        header=0 if has_header else None,
        on_bad_lines="error",
        low_memory=False,
    )
    try:
        df = pd.read_csv(path, **read_kwargs)
    except UnicodeDecodeError:
        warnings.append("UTF-8 decoding failed mid-file; re-read as latin-1.")
        read_kwargs["encoding"] = "latin-1"
        try:
            df = pd.read_csv(path, **read_kwargs)
        except pd.errors.ParserError as exc:
            raise IngestionError("MALFORMED_CSV", f"Malformed CSV rows: {exc}")
    except pd.errors.EmptyDataError:
        raise IngestionError("MALFORMED_CSV", "CSV contains no parseable columns.")
    except pd.errors.ParserError as exc:
        raise IngestionError(
            "MALFORMED_CSV",
            f"Malformed CSV rows (ragged or unbalanced quotes): {exc}",
        )

    if df.shape[1] == 0:
        raise IngestionError("MALFORMED_CSV", "CSV parsed to zero columns.")

    if not has_header:
        df.columns = [f"col_{i}" for i in range(df.shape[1])]

    # Duplicate header names: pandas already mangles them; surface a warning.
    if len(set(df.columns)) != len(df.columns):
        warnings.append("Duplicate column names detected; pandas renamed them (x.1).")

    if df.empty:
        warnings.append("CSV has a header but zero data rows.")

    dest = settings.DATA_DIR / "staging" / f"{uuid.uuid4().hex}.parquet"
    table = _df_to_table(df, dest, logical_name)
    table.warnings.extend(warnings)
    return table


# ---------------------------------------------------------------------------
# Excel (.xlsx / .xls)
# ---------------------------------------------------------------------------

def _excel_sheet_meta(path: Path) -> Dict[str, Any]:
    """Workbook / sheet names / dimensions WITHOUT loading cell data (xlsx)."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = []
        for name in wb.sheetnames:
            ws = wb[name]
            rows = None
            cols = None
            try:
                if ws.max_row:
                    rows = int(ws.max_row)
                if ws.max_column:
                    cols = int(ws.max_column)
            except Exception:
                pass
            sheets.append({"name": name, "max_rows": rows, "max_columns": cols})
        return {"engine": "openpyxl", "sheets": sheets}
    finally:
        wb.close()


def parse_excel(path: Path, logical_name: str, sheet_name: Optional[str]) -> TableInfo:
    warnings: List[str] = []
    ext = path.suffix.lower()

    if ext == ".xlsx":
        meta = _excel_sheet_meta(path)
        sheets = [s["name"] for s in meta["sheets"]]
        if sheet_name:
            if sheet_name not in sheets:
                raise IngestionError(
                    "SHEET_NOT_FOUND",
                    f"Sheet {sheet_name!r} not found. Available sheets: {sheets}",
                )
            chosen = sheet_name
        else:
            chosen = sheets[0]
            if len(sheets) > 1:
                warnings.append(
                    f"Workbook has {len(sheets)} sheets {sheets}; ingested first "
                    f"sheet {chosen!r}. Re-upload with sheet_name to select another."
                )
        try:
            df = pd.read_excel(path, sheet_name=chosen, engine="openpyxl")
        except Exception as exc:  # malformed workbook / broken XML
            raise IngestionError(
                "MALFORMED_EXCEL",
                f"Could not read workbook sheet {chosen!r}: {type(exc).__name__}: {exc}",
            )
        meta_sheets = meta["sheets"]
    else:  # .xls via xlrd (read-only, small dependency)
        chosen = None
        try:
            import xlrd

            book = xlrd.open_workbook(path, on_demand=True)
            names = book.sheet_names()
            if sheet_name:
                if sheet_name not in names:
                    raise IngestionError(
                        "SHEET_NOT_FOUND",
                        f"Sheet {sheet_name!r} not found. Available sheets: {names}",
                    )
                chosen = sheet_name
            else:
                chosen = names[0] if names else None
                if len(names) > 1:
                    warnings.append(
                        f"Workbook has {len(names)} sheets {names}; ingested first "
                        f"sheet {chosen!r}. Re-upload with sheet_name to select another."
                    )
            book.release_resources()
        except IngestionError:
            raise
        except Exception as exc:
            raise IngestionError(
                "MALFORMED_EXCEL",
                f"Could not read legacy .xls workbook: {type(exc).__name__}: {exc}",
            )
        try:
            df = pd.read_excel(path, sheet_name=chosen, engine="xlrd")
        except Exception as exc:
            raise IngestionError(
                "MALFORMED_EXCEL",
                f"Could not read sheet {chosen!r}: {type(exc).__name__}: {exc}",
            )
        meta_sheets = [{"name": n, "max_rows": None, "max_columns": None} for n in names]

    if df.empty and df.shape[1] == 0:
        raise IngestionError("MALFORMED_EXCEL", "Selected sheet is empty.")

    table_name = f"{logical_name}::{chosen}" if chosen else logical_name
    dest = settings.DATA_DIR / "staging" / f"{uuid.uuid4().hex}.parquet"
    table = _df_to_table(df, dest, table_name)
    table.warnings.extend(warnings)
    table_warnings = table.warnings
    return table


# ---------------------------------------------------------------------------
# JSON (array / NDJSON)
# ---------------------------------------------------------------------------

def _flatten_records(records: List[Dict[str, Any]], warnings: List[str]) -> List[Dict[str, Any]]:
    if any(isinstance(v, (dict, list)) for r in records for v in r.values()):
        warnings.append("Nested JSON structures were flattened with dot notation.")
        return pd.json_normalize(records).to_dict(orient="records")
    return records


def parse_json(path: Path, logical_name: str) -> TableInfo:
    warnings: List[str] = []
    raw = path.read_bytes()
    if not raw.strip():
        raise IngestionError("EMPTY_FILE", "JSON file is empty.")
    text = raw.decode("utf-8-sig", errors="replace")

    records: Optional[List[Dict[str, Any]]] = None

    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            if not parsed:
                raise IngestionError("EMPTY_JSON", "JSON array is empty.")
            if all(isinstance(x, dict) for x in parsed):
                records = parsed
            elif all(isinstance(x, (int, float, str, bool)) for x in parsed):
                records = [{"value": x} for x in parsed]
                warnings.append("Scalar JSON array normalized to a single 'value' column.")
            elif all(isinstance(x, list) for x in parsed):
                raise IngestionError(
                    "UNSUPPORTED_JSON_STRUCTURE",
                    "Top-level array of arrays is not supported; expected an array "
                    "of objects or newline-delimited JSON.",
                )
            else:
                raise IngestionError(
                    "UNSUPPORTED_JSON_STRUCTURE",
                    "Mixed-type JSON array is not supported; expected uniform objects.",
                )
        elif isinstance(parsed, dict):
            # Single JSON object: either nested doc or one NDJSON record.
            if any(isinstance(v, list) for v in parsed.values()):
                raise IngestionError(
                    "UNSUPPORTED_JSON_STRUCTURE",
                    "Top-level object with array values is not supported; provide a "
                    "JSON array of objects or newline-delimited JSON.",
                )
            records = [parsed]
            warnings.append("Single JSON object treated as a one-row dataset.")
        else:
            raise IngestionError(
                "UNSUPPORTED_JSON_STRUCTURE",
                f"Unsupported top-level JSON structure: {type(parsed).__name__}.",
            )
    except json.JSONDecodeError as exc:
        # Fallback: newline-delimited JSON.
        records_nd: List[Dict[str, Any]] = []
        for i, line in enumerate(text.splitlines(), start=1):
            line = line.strip()
            if not line:
                continue
            try:
                obj = json.loads(line)
            except json.JSONDecodeError as jexc:
                raise IngestionError(
                    "MALFORMED_JSON",
                    f"Invalid JSON at line {i}, column {jexc.colno}: {jexc.msg}",
                )
            if not isinstance(obj, dict):
                raise IngestionError(
                    "UNSUPPORTED_JSON_STRUCTURE",
                    f"NDJSON line {i} is {type(obj).__name__}; expected an object per line.",
                )
            records_nd.append(obj)
        if not records_nd:
            raise IngestionError(
                "MALFORMED_JSON",
                f"Not valid JSON (line {exc.lineno}, column {exc.colno}: {exc.msg}) "
                "and no newline-delimited records found.",
            )
        records = records_nd
        warnings.append("Parsed as newline-delimited JSON (NDJSON).")

    assert records is not None
    records = _flatten_records(records, warnings)
    df = pd.DataFrame.from_records(records)
    if df.empty:
        raise IngestionError("EMPTY_JSON", "JSON produced zero rows.")

    dest = settings.DATA_DIR / "staging" / f"{uuid.uuid4().hex}.parquet"
    table = _df_to_table(df, dest, logical_name)
    table.warnings.extend(warnings)
    return table


# ---------------------------------------------------------------------------
# Parquet (lazy/columnar - original file is already normalized)
# ---------------------------------------------------------------------------

def parse_parquet(path: Path, logical_name: str) -> TableInfo:
    warnings: List[str] = []
    try:
        schema, rows = _schema_from_parquet(path)
    except Exception as exc:
        raise IngestionError(
            "MALFORMED_PARQUET",
            f"Not a valid Parquet file: {type(exc).__name__}: {exc}",
        )
    if not schema:
        raise IngestionError("MALFORMED_PARQUET", "Parquet file declares zero columns.")
    return TableInfo(
        table_name=logical_name,
        normalized_path=str(path),  # already columnar; no conversion
        row_count=rows,
        column_count=len(schema),
        schema=schema,
        warnings=warnings,
    )


# ---------------------------------------------------------------------------
# SQL (untrusted; controlled allow-list execution layer)
# ---------------------------------------------------------------------------

_SQL_FORBIDDEN = re.compile(
    r"\b(DROP|DELETE|UPDATE|ALTER|TRUNCATE|ATTACH|DETACH|LOAD|INSTALL|COPY|"
    r"EXPORT|IMPORT|CALL|PRAGMA|GRANT|REVOKE|VACUUM|BEGIN|COMMIT|ROLLBACK|"
    r"SET|RESET|SHUTDOWN|EXECUTE|EXEC|MERGE|REPLACE|CREATE\s+(INDEX|VIEW|"
    r"TRIGGER|SCHEMA|SEQUENCE|DATABASE)|INTO\s+OUTFILE)\b",
    re.IGNORECASE,
)
_SQL_READ_FUNCS = re.compile(
    r"\b(read_csv|read_parquet|read_json|read_text|glob|parquet_scan|"
    r"csv_scan|sniff_csv|delta_scan|iceberg_scan)\s*\(",
    re.IGNORECASE,
)
_CREATE_TABLE_RE = re.compile(
    r"CREATE\s+TABLE\s+(?:IF\s+NOT\s+EXISTS\s+)?"
    r"(?:\"([^\"]+)\"|`([^`]+)`|\[([^\]]+)\]|([A-Za-z_][A-Za-z0-9_]*))",
    re.IGNORECASE,
)


def split_sql_statements(text: str) -> List[str]:
    """Split on ';' while respecting quotes and comments."""
    statements: List[str] = []
    buf: List[str] = []
    i = 0
    n = len(text)
    in_single = in_double = in_line = in_block = False
    while i < n:
        ch = text[i]
        nxt = text[i + 1] if i + 1 < n else ""
        if in_line:
            buf.append(ch)
            if ch == "\n":
                in_line = False
        elif in_block:
            buf.append(ch)
            if ch == "*" and nxt == "/":
                buf.append(nxt)
                i += 1
                in_block = False
        elif in_single:
            buf.append(ch)
            if ch == "'":
                if nxt == "'":
                    buf.append(nxt)
                    i += 1
                else:
                    in_single = False
        elif in_double:
            buf.append(ch)
            if ch == '"':
                if nxt == '"':
                    buf.append(nxt)
                    i += 1
                else:
                    in_double = False
        else:
            if ch == "-" and nxt == "-":
                in_line = True
                buf.append(ch)
            elif ch == "/" and nxt == "*":
                in_block = True
                buf.append(ch)
            elif ch == "'":
                in_single = True
                buf.append(ch)
            elif ch == '"':
                in_double = True
                buf.append(ch)
            elif ch == ";":
                statements.append("".join(buf))
                buf = []
            else:
                buf.append(ch)
        i += 1
    statements.append("".join(buf))
    return [s for s in ("".join([st]).strip() for st in statements) if s]


def _strip_sql_comments(stmt: str) -> str:
    stmt = re.sub(r"--[^\n]*", " ", stmt)
    stmt = re.sub(r"/\*.*?\*/", " ", stmt, flags=re.DOTALL)
    return stmt.strip()


def _sanitize_table_name(name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_]+", "_", name).strip("_").lower()
    return name[:60] or "table"


def parse_sql(path: Path, logical_name: str) -> List[TableInfo]:
    """Load CREATE TABLE / INSERT INTO (VALUES) dumps into an in-memory DuckDB.

    Only the allow-listed statements are ever executed, in a throwaway
    database, with filesystem-capable functions rejected up front.
    """
    warnings: List[str] = []
    raw = path.read_bytes()
    if not raw.strip():
        raise IngestionError("EMPTY_FILE", "SQL file is empty.")
    if b"\x00" in raw[:8192]:
        raise IngestionError("BINARY_NOT_TEXT", "SQL file contains binary data.")
    text = raw.decode("utf-8-sig", errors="replace")

    statements = [s for s in (_strip_sql_comments(s) for s in split_sql_statements(text)) if s]
    if not statements:
        raise IngestionError("EMPTY_SQL", "SQL file contains no statements.")

    allowed: List[str] = []
    rejected: List[dict] = []
    for stmt in statements:
        m = _SQL_FORBIDDEN.search(stmt) or _SQL_READ_FUNCS.search(stmt)
        if m:
            rejected.append(
                {
                    "statement_start": stmt[:80],
                    "reason": f"Forbidden construct {m.group(0)!r} rejected.",
                }
            )
            continue
        if re.match(r"CREATE\s+TABLE\b", stmt, re.IGNORECASE):
            if re.search(r"\bAS\s+SELECT\b", stmt, re.IGNORECASE):
                rejected.append(
                    {
                        "statement_start": stmt[:80],
                        "reason": "CREATE TABLE ... AS SELECT is not permitted.",
                    }
                )
                continue
            allowed.append(stmt)
            continue
        if re.match(r"INSERT\s+INTO\b", stmt, re.IGNORECASE):
            if re.search(r"\bSELECT\b", stmt, re.IGNORECASE):
                rejected.append(
                    {
                        "statement_start": stmt[:80],
                        "reason": "INSERT ... SELECT is not permitted (file access).",
                    }
                )
                continue
            if not re.search(r"\bVALUES\b", stmt, re.IGNORECASE):
                rejected.append(
                    {
                        "statement_start": stmt[:80],
                        "reason": "INSERT without VALUES is not permitted.",
                    }
                )
                continue
            allowed.append(stmt)
            continue
        rejected.append(
            {
                "statement_start": stmt[:80],
                "reason": "Unsupported statement; only CREATE TABLE and "
                "INSERT INTO ... VALUES are accepted.",
            }
        )

    if not allowed:
        reasons = "; ".join(r["reason"] for r in rejected[:5])
        raise IngestionError(
            "SQL_NO_SUPPORTED_STATEMENTS",
            "No supported SQL statements found (only CREATE TABLE and "
            f"INSERT INTO ... VALUES are accepted). Rejected: {reasons}",
        )

    import duckdb

    conn = duckdb.connect(database=":memory:")
    tables: List[TableInfo] = []
    created_tables: List[str] = []
    try:
        # Phase 1: execute every allowed statement first...
        for stmt in allowed:
            try:
                conn.execute(stmt)
            except Exception as exc:
                raise IngestionError(
                    "SQL_EXECUTION_FAILED",
                    f"Statement rejected by engine: {stmt[:80]!r} -> "
                    f"{type(exc).__name__}: {exc}",
                )
            m = _CREATE_TABLE_RE.match(_strip_sql_comments(stmt))
            if m:
                raw_name = next(g for g in m.groups() if g)
                created_tables.append(_sanitize_table_name(raw_name))

        # Phase 2: ...then materialize each table (now fully populated).
        for name in created_tables:
            dest = (
                settings.DATA_DIR
                / "staging"
                / f"{uuid.uuid4().hex}__{name}.parquet"
            )
            dest.parent.mkdir(parents=True, exist_ok=True)
            ident = name.replace('"', '""')
            path_lit = str(dest).replace("'", "''")
            try:
                conn.execute(f'COPY (SELECT * FROM "{ident}") TO \'{path_lit}\'')
            except Exception as exc:
                raise IngestionError(
                    "SQL_EXPORT_FAILED",
                    f"Could not materialize table {name!r}: {type(exc).__name__}: {exc}",
                )
            schema, rows = _schema_from_parquet(dest)
            if rows == 0:
                warnings.append(f"SQL table {name!r} was created but has zero rows.")
            tables.append(
                TableInfo(
                    table_name=name,
                    normalized_path=str(dest),
                    row_count=rows,
                    column_count=len(schema),
                    schema=schema,
                    warnings=list(warnings),
                )
            )
    finally:
        conn.close()

    if not tables:
        raise IngestionError(
            "SQL_NO_TABLE_CREATED", "SQL file did not create any tables."
        )
    if rejected:
        for t in tables:
            t.warnings.append(
                f"{len(rejected)} statement(s) rejected by the safe SQL layer "
                "(only CREATE TABLE / INSERT INTO ... VALUES allowed)."
            )
    return tables


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------

def parse_file(
    path: Path, logical_name: str, sheet_name: Optional[str] = None
) -> List[TableInfo]:
    ext = path.suffix.lower()
    if ext == ".csv":
        return [parse_csv(path, logical_name)]
    if ext in (".xlsx", ".xls"):
        return [parse_excel(path, logical_name, sheet_name)]
    if ext == ".json":
        return [parse_json(path, logical_name)]
    if ext == ".parquet":
        return [parse_parquet(path, logical_name)]
    if ext == ".sql":
        return parse_sql(path, logical_name)
    raise IngestionError(
        "UNSUPPORTED_EXTENSION",
        f"Unsupported file type {ext!r}. Supported: csv, xlsx, xls, json, "
        "parquet, sql (+ zip archives).",
        http_status=415,
    )
